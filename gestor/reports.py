from openpyxl import Workbook
from openpyxl.styles import Font, Border, Alignment, Side, PatternFill
from openpyxl.utils import get_column_letter
from django.utils import timezone
from pathlib import Path
from django.conf import settings
from .models import Producto
class reporteEXCEL:
    """Exportar una lista de tuplas a una tabla en un archivo xlsx (EXCEL)."""

    def __init__(self, titulo, cabecera, registros, nombreEXCEL):
        self.titulo = titulo
        self.cabecera = cabecera
        self.registros = registros
        self.nombreEXCEL = nombreEXCEL

    def Exportar(self):
        libroTrabajo = Workbook()
        hoja = libroTrabajo.active
        hoja.title = self.titulo
        hoja.sheet_properties.tabColor = "1072BA"
        hoja.sheet_view.showGridLines = False

        # ====== Cálculo correcto de la última columna ======
        # Encabezado comienza en columna 2 (B) y hay len(self.cabecera) columnas,
        # la última columna es 1 + len(self.cabecera)
        last_col_idx = 1 + len(self.cabecera)
        celdaFinal = get_column_letter(last_col_idx)

        rangoTitulo = f"B2:{celdaFinal}3"
        rangoCabecera = f"B10:{celdaFinal}10"

        centrarTexto = Alignment(horizontal="center", vertical="center")

        # ========================== TÍTULO ==========================
        hoja.merge_cells(rangoTitulo)
        celdaTitulo = hoja.cell(row=2, column=2)
        celdaTitulo.value = self.titulo.upper()
        celdaTitulo.alignment = centrarTexto
        celdaTitulo.font = Font(color="FF000000", size=11, bold=True)

        # ===================== INFORMACIÓN EXTRA ====================
        fontInformacionExtra = Font(color="707070", size=11, bold=False)

        celdaOrigen = hoja.cell(row=5, column=2)
        
        celdaOrigen.font = fontInformacionExtra

        celdaFechaDescarga = hoja.cell(row=6, column=2)
        # Si no usas Django, reemplaza por: datetime.now().strftime("%d/%m/%Y")
        celdaFechaDescarga.value = f"Fecha de descarga: {timezone.localtime().strftime('%d/%m/%Y')}"
        celdaFechaDescarga.font = fontInformacionExtra

        celdaCantidadDescarga = hoja.cell(row=8, column=2)
        celdaCantidadDescarga.value = f"Registros descargados: {len(self.registros)}"
        celdaCantidadDescarga.font = fontInformacionExtra

        # ================== BORDES/COLORES ==========================
        thin = Side(border_style="thin", color="000000")
        border = Border(top=thin, left=thin, right=thin, bottom=thin)
        colorCelda = PatternFill("solid", fgColor="C0C0C0")

        # Borde + relleno en TÍTULO
        filasTitulo = hoja[rangoTitulo]
        celdaInicial = filasTitulo[0][0].row
        for fila in filasTitulo:
            filaIzquierda = fila[0]
            filaDerecha = fila[-1]
            filaIzquierda.border = filaIzquierda.border + Border(left=border.left)
            filaDerecha.border = filaDerecha.border + Border(right=border.right)
            for celda in fila:
                if celda.row == celdaInicial:
                    celda.border = celda.border + Border(top=border.top)
                else:
                    celda.border = celda.border + Border(bottom=border.bottom)
                celda.fill = colorCelda

        # CABECERA de la tabla
        for idx, texto in enumerate(self.cabecera, start=2):
            c = hoja.cell(row=10, column=idx)
            c.value = texto
            c.border = border
            c.alignment = centrarTexto
            c.font = Font(color="FF000000", size=10, bold=True)

        for fila in hoja[rangoCabecera]:
            for celda in fila:
                celda.fill = colorCelda

        # REGISTROS
        for filaIndice, fila_regs in enumerate(self.registros, start=11):
            for colIndice, valor in enumerate(fila_regs, start=2):
                c = hoja.cell(row=filaIndice, column=colIndice)
                c.value = valor
                c.border = border
                c.alignment = Alignment(horizontal="left", vertical="center")
                c.font = Font(color="FF000000", size=10, bold=False)

        # ============== AUTO-AJUSTE DE ANCHO POR COLUMNA ===========
        # Recorremos SOLO las columnas reales: B .. celdaFinal
        for col_idx in range(2, last_col_idx + 1):
            col_letter = get_column_letter(col_idx)
            max_len = 0

            # incluir header (fila 10) y todos los datos hasta el final
            for row in range(10, hoja.max_row + 1):
                val = hoja.cell(row=row, column=col_idx).value
                if val is not None:
                    max_len = max(max_len, len(str(val)))

            # margen extra para que no quede muy justo
            hoja.column_dimensions[col_letter].width = max(8, max_len + 2)

        try:
            libroTrabajo.save(f"{self.nombreEXCEL}.xlsx")
            retornar = "Reporte generado con éxito."
        except PermissionError:
            retornar = "Error inesperado: Permiso denegado (archivo abierto o sin permisos)."
        except Exception as e:
            retornar = f"Error desconocido: {e}"
        finally:
            libroTrabajo.close()
            return retornar


def generarReporte():
    from .models import Producto, reporteEXCEL  # asegúrate de importar tu modelo y clase Excel aquí

    # Título y cabecera del reporte
    titulo = "LISTADO DE PRODUCTOS"
    cabecera = ("Nombre", "Descripción", "Cantidad", "Precio", "Categoría", "Fecha de Entrada")

    # Consultar los productos existentes
    productos = Producto.objects.all().order_by('-fecha_entrada')

    # Convertir queryset en una lista de tuplas
    registros = [
        (
            p.nombre,
            p.descripcion,
            p.cantidad,
            float(p.precio),
            p.categoria,
            p.fecha_entrada.strftime("%d/%m/%Y"),
        )
        for p in productos
    ]

    # Carpeta destino para guardar el archivo
    carpeta = Path(settings.MEDIA_ROOT) / "reportes"
    carpeta.mkdir(parents=True, exist_ok=True)

    # Nombre único basado en la fecha/hora
    stamp = timezone.localtime().strftime('%Y%m%d_%H%M%S')
    nombre_sin_ext = carpeta / f"reporte_productos_{stamp}"

    # Crear el Excel
    mensaje = reporteEXCEL(titulo, cabecera, registros, str(nombre_sin_ext)).Exportar()

    # URL pública para el template
    url_publica = f"{settings.MEDIA_URL}reportes/{nombre_sin_ext.name}.xlsx"
    return mensaje, url_publica